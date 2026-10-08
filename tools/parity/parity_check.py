#!/usr/bin/env python3
"""
Paridad funcional: aplicación de REFERENCIA (HTML original) vs build nuevo. Ejecuta el mismo guion en ambas (carga de SAC/plantillas, archivo dañado,
perfil, navegación con contexto, vista previa DOCX con variables, tema, atajos, mantenimiento, recarga/persistencia) y compara. Sale con código ≠ 0 si difieren.
Requisitos (solo desarrollo): pip install playwright openpyxl python-docx && playwright install chromium
Uso: npm run build && npm run preview -- --port 4173   (otra terminal)   →   python tools/parity/parity_check.py --url http://localhost:4173/
Opcional: CHROMIUM_PATH=/ruta/al/chrome
"""
import argparse, os, pathlib, sys, tempfile, datetime, random
ROOT = pathlib.Path(__file__).resolve().parents[2]
DEFAULT_REF = ROOT / 'reference' / 'Asistente Documental.original.html'

def make_data(d: pathlib.Path):
    import openpyxl, docx
    cols = ['NUMERO_CUENTA','NUMERO_PROCESO','PROCESO','FECHA_SOLICITUD','FECHA_VENCIMIENTO','MEDIO_SOLICITUD','NOMBRE_USUARIO_INICIAL_PROCESO','NOMBRE_SOLICITANTE','MUNICIPIO_SOLICITANTE','RADICADO_ENTRADA','RADICADO_SALIDA','PRIMER_NOMBRE']
    random.seed(7); wb = openpyxl.Workbook(); ws = wb.active; ws.append(cols); base = datetime.datetime(2026, 10, 8)
    for i in range(120):
        f = base - datetime.timedelta(days=random.randint(0, 40))
        ws.append([str(1000000+i), 'P%07d'%i, 'PQR', f, f+datetime.timedelta(days=15), 'Escrito', 'USUARIO DE PRUEBA', 'Ana Gómez', 'Girón', '2026%010d'%(i+1000), '', 'Ana'])
    wb.save(d/'sac.xlsx'); (d/'corrupt.xlsx').write_bytes(b'PK no es un xlsx real '*40); (d/'tpl').mkdir(exist_ok=True)
    for name in ('Respuesta general.docx','Respuesta factura.docx'):
        doc = docx.Document(); doc.add_heading('ESSA E.S.P.', 1)
        for line in ('Señor(a) [PRIMER_NOMBRE]','Radicado [RADICADO_ENTRADA]','Municipio [MUNICIPIO_SOLICITANTE]'): doc.add_paragraph(line)
        doc.save(d/'tpl'/name)

SNAP = '''()=>{const s=AsistenteDocumental.store.getState();const out={};
 for(const k of s.order){const r=s.resources[k];out[k]={status:r.status,rows:r.rowCount,err:r.error,warn:r.warnings,file:r.file&&{name:r.file.name,count:r.file.count,items:r.file.items&&r.file.items.map(i=>i.fileName)},persisted:r.persisted}}
 return {res:out,hydrated:s.hydrated,tabs:[...document.querySelectorAll('.ad-tab')].map(t=>t.getAttribute('aria-label')+'|'+t.dataset.state)}}'''

def scenario(pw, url, data, chromium):
    b = pw.chromium.launch(executable_path=chromium, args=['--no-sandbox']) if chromium else pw.chromium.launch(args=['--no-sandbox'])
    pg = b.new_context(viewport={'width':1440,'height':900}).new_page(); errs, out = [], {}
    pg.on('pageerror', lambda e: errs.append('PAGEERR '+str(e))); pg.on('console', lambda m: errs.append('CONSOLE '+m.text) if m.type=='error' else None)
    pg.goto(url); pg.wait_for_timeout(2500)
    out['api'] = pg.evaluate('''()=>({keys:Object.keys(AsistenteDocumental),version:AsistenteDocumental.version,libs:Object.entries(AsistenteDocumental.libs).map(([k,v])=>k+':'+typeof v),
      globals:['JSZip','PizZip','docxtemplater','easyTemplateX','docx','docxRenderer','mammoth','ADDocxLocal','XLSXLite','PersonName','DocStore','DataRepository','BackupStore'].map(k=>k+':'+typeof window[k]),frames:[...document.querySelectorAll('iframe')].map(f=>f.id)})''')
    rec = pg.frame_locator('#panel-recursos')
    rec.locator('section[data-key=sac] input[type=file]').set_input_files(str(data/'corrupt.xlsx')); pg.wait_for_timeout(1500); out['corrupt'] = pg.evaluate(SNAP)
    rec.locator('section[data-key=sac] input[type=file]').set_input_files(str(data/'sac.xlsx')); pg.wait_for_timeout(2500)
    rec.locator('section[data-key=plantillas] input[webkitdirectory]').set_input_files(str(data/'tpl')); pg.wait_for_timeout(2500)
    out['loaded'] = pg.evaluate(SNAP)
    out['rows'] = pg.evaluate("()=>{const r=AsistenteDocumental.store.getRows('sac');return JSON.parse(JSON.stringify({n:r.length,first:r[0],last:r[r.length-1]}))}")
    pg.evaluate("AsistenteDocumental.store.setProfile({nombre:'Atención Clientes',cargo:'Analista',firma:null})"); pg.wait_for_timeout(400)
    rad = pg.evaluate("AsistenteDocumental.store.getRows('sac')[0].RADICADO_ENTRADA")
    out['nav'] = pg.evaluate("rad=>AsistenteDocumental.navigate('documentos',{radicado:String(rad)},{steps:['A','B','C']})", rad); pg.wait_for_timeout(3000)
    do = pg.frame_locator('#panel-documentos'); do.locator('#plantillaList > *').nth(1).click(); pg.wait_for_timeout(3000)
    out['preview'] = do.locator('#previewCanvas').inner_text()[:600]
    pg.click('#adTheme'); pg.wait_for_timeout(500)
    out['theme'] = pg.evaluate("[document.documentElement.dataset.theme,localStorage.getItem('essa-theme'),[...document.querySelectorAll('iframe')].map(f=>f.contentDocument.documentElement.dataset.theme).join()]")
    for i in (1,2,3):
        pg.keyboard.press(f'Alt+{i}'); pg.wait_for_timeout(900); out[f'alt{i}'] = pg.evaluate("[location.hash,document.querySelector('.ad-frame.is-active').id]")
    out['nav_bad'] = pg.evaluate("AsistenteDocumental.navigate('nope')")
    out['nav_fail'] = pg.evaluate("AsistenteDocumental.navigate('documentos',{radicado:'999'},{steps:['A','B']})")
    out['maintenance'] = pg.evaluate("AsistenteDocumental.runMaintenance().then(r=>JSON.stringify(r))")
    pg.goto(url); pg.wait_for_timeout(4000); out['reload'] = pg.evaluate(SNAP); out['errors'] = errs; b.close(); return out

def diff(a, b, path=''):
    if type(a) != type(b): return [f'tipo distinto en {path}']
    if isinstance(a, dict):
        r = []
        for k in sorted(set(a)|set(b)): r += [f'falta {path}/{k}'] if (k not in a or k not in b) else diff(a[k], b[k], f'{path}/{k}')
        return r
    if isinstance(a, list):
        if len(a) != len(b): return [f'longitud distinta en {path}: {len(a)} vs {len(b)}']
        return [d for i,(x,y) in enumerate(zip(a,b)) for d in diff(x,y,f'{path}[{i}]')]
    return [] if a == b else [f'{path}: {str(a)[:120]!r} ≠ {str(b)[:120]!r}']

def main():
    ap = argparse.ArgumentParser(); ap.add_argument('--url', required=True); ap.add_argument('--reference', default=str(DEFAULT_REF)); a = ap.parse_args()
    from playwright.sync_api import sync_playwright
    data = pathlib.Path(tempfile.mkdtemp(prefix='ad-parity-')); make_data(data); chromium = os.environ.get('CHROMIUM_PATH')
    with sync_playwright() as pw:
        ref = scenario(pw, pathlib.Path(a.reference).resolve().as_uri(), data, chromium); new = scenario(pw, a.url, data, chromium)
    problems = diff(ref, new)
    if ref['errors'] or new['errors']: problems.append(f"errores de consola: referencia={ref['errors']} nuevo={new['errors']}")
    print('PARIDAD OK: la versión nueva se comporta igual que la referencia.' if not problems else 'DIFERENCIAS:\n - '+'\n - '.join(problems))
    print('Pasos comparados:', ', '.join(k for k in ref if k != 'errors')); sys.exit(1 if problems else 0)
if __name__ == '__main__': main()
